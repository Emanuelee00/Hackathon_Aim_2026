import asyncio
import json
from io import BytesIO
from unittest.mock import AsyncMock, MagicMock

import httpx
from docx import Document

from employment.extraction import extract_cv_text
from employment.models import EmploymentContext
from employment.planning import generate_plan
from employment.routes import employment_plan


def context() -> EmploymentContext:
    return EmploymentContext(
        resident_name="Fatou",
        objective="Commis de cuisine",
        acquired_skills=["Cuisine collective"],
    )


def test_extracts_text_from_docx():
    document = Document()
    document.add_paragraph("Fatou K. — Expérience en cuisine collective")
    document.add_paragraph("Préparation de repas et organisation d'équipe")
    output = BytesIO()
    document.save(output)
    text = extract_cv_text("cv.docx", output.getvalue())
    assert "cuisine collective" in text
    assert "Préparation de repas" in text


def test_employment_plan_uses_valid_ai_response(monkeypatch):
    result = {
        "summary": "Profil cohérent avec la restauration.",
        "strengths": ["Cuisine collective"],
        "gaps": ["Préciser les volumes préparés"],
        "cv_suggestions": ["Ajouter l'expérience Chez Marthe"],
        "steps": [
            {
                "title": "CV",
                "action": "Adapter le titre.",
                "timeframe": "Cette semaine",
            },
            {
                "title": "Offres",
                "action": "Choisir trois offres.",
                "timeframe": "Sous 10 jours",
            },
            {
                "title": "Candidature",
                "action": "Envoyer un CV adapté.",
                "timeframe": "Sous 2 semaines",
            },
        ],
    }
    transport = httpx.MockTransport(
        lambda request: httpx.Response(
            200, json={"message": {"content": json.dumps(result)}}
        )
    )
    monkeypatch.setattr(
        "employment.planning.httpx.Client",
        MagicMock(return_value=httpx.Client(transport=transport)),
    )
    plan = generate_plan(
        context(), "Expérience longue en cuisine collective et préparation de repas."
    )
    assert plan.source == "ai"
    assert plan.steps[0].title == "CV"


def test_employment_plan_falls_back_when_model_is_unavailable(monkeypatch):
    monkeypatch.setattr(
        "employment.planning.httpx.Client",
        MagicMock(side_effect=httpx.ConnectError("offline")),
    )
    plan = generate_plan(
        context(),
        "Expérience en accueil, organisation et préparation de repas collectifs.",
    )
    assert plan.source == "guided"
    assert len(plan.steps) == 3
    assert plan.strengths == ["Cuisine collective"]


def test_employment_endpoint_accepts_a_real_docx(monkeypatch):
    document = Document()
    document.add_paragraph(
        "Expérience en restauration collective et préparation de repas"
    )
    document.add_paragraph("Accueil des participantes et organisation du matériel")
    output = BytesIO()
    document.save(output)
    monkeypatch.setattr(
        "employment.planning.httpx.Client",
        MagicMock(side_effect=httpx.ConnectError("offline")),
    )
    upload = MagicMock(filename="cv.docx")
    upload.read = AsyncMock(return_value=output.getvalue())
    plan = asyncio.run(
        employment_plan(
            cv=upload,
            objective="Commis de cuisine",
            resident_name="Fatou",
            acquired_skills='["Cuisine collective"]',
        )
    )
    assert plan.source == "guided"
